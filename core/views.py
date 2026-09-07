import openpyxl
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from django.contrib import messages
from django.utils import timezone
from .models import (
    SchoolInfo, ClassSection, Subject, Student, Employee, 
    BankDetail, FeeParticular, FeeStructure, FeePayment, 
    Attendance, EmployeeAttendance, Exam, ExamMark, Homework, MessageNotice
)

def dashboard(request):
    context = {
        'total_students': Student.objects.count(),
        'total_employee': Employee.objects.count(),
        'total_classes': ClassSection.objects.count(),
        'school_info': SchoolInfo.objects.first(),
    }
    return render(request, 'core/dashboard.html', context)

def add_student(request):
    if request.method == 'POST':
        reg_no = request.POST.get('registration_no')
        first_name = request.POST.get('first_name')
        last_name = request.POST.get('last_name', '')
        cls_name = request.POST.get('class_name')
        section = request.POST.get('section', '')
        
        class_obj, _ = ClassSection.objects.get_or_create(class_name=cls_name, section=section)
        
        Student.objects.create(
            registration_no=reg_no,
            student_name=f"{first_name} {last_name}".strip(),
            first_name=first_name,
            last_name=last_name,
            class_section=class_obj,
            section=section,
            roll_no=request.POST.get('roll_no') or None,
            gender=request.POST.get('gender'),
            dob=request.POST.get('dob') or None,
            father_name=request.POST.get('father_name'),
            mother_name=request.POST.get('mother_name'),
            phone_number=request.POST.get('phone_number'),
            address=request.POST.get('address')
        )
        messages.success(request, "Student admitted successfully!")
        return redirect('dashboard')
    
    classes = ClassSection.objects.all()
    return render(request, 'core/add_student.html', {'classes': classes})

def add_new_class(request):
    if request.method == 'POST':
        c_name = request.POST.get('class_name')
        sec = request.POST.get('section', '')
        ClassSection.objects.get_or_create(class_name=c_name, section=sec)
        messages.success(request, "Class added successfully!")
        return redirect('classes_list')
    classes = ClassSection.objects.all()
    return render(request, 'core/classes.html', {'classes': classes})

def classes_list(request):
    classes = ClassSection.objects.all()
    return render(request, 'core/classes.html', {'classes': classes})

def subjects_view(request):
    if request.method == 'POST':
        cls_id = request.POST.get('class_section')
        sub_name = request.POST.get('subject_name')
        sub_code = request.POST.get('subject_code')
        cls_obj = get_object_or_404(ClassSection, id=cls_id)
        Subject.objects.create(class_section=cls_obj, subject_name=sub_name, subject_code=sub_code)
        messages.success(request, "Subject added successfully!")
        return redirect('subjects')
    subjects = Subject.objects.all()
    classes = ClassSection.objects.all()
    return render(request, 'core/subjects.html', {'subjects': subjects, 'classes': classes})

def school_info_view(request):
    school, _ = SchoolInfo.objects.get_or_create(id=1, defaults={'name': 'Guru Nanak Sr Sec School'})
    if request.method == 'POST':
        school.name = request.POST.get('name', school.name)
        school.phone = request.POST.get('phone', school.phone)
        school.email = request.POST.get('email', school.email)
        school.address = request.POST.get('address', school.address)
        school.save()
        messages.success(request, "School Info updated successfully!")
        return redirect('school_info')
    return render(request, 'core/schoolinfo.html', {'school': school})

def submit_fees_view(request):
    if request.method == 'POST':
        student_id = request.POST.get('student')
        particular_id = request.POST.get('particular')
        amount = request.POST.get('amount_paid')
        receipt_no = request.POST.get('receipt_no')
        
        student = get_object_or_404(Student, id=student_id)
        particular = get_object_or_404(FeeParticular, id=particular_id)
        
        FeePayment.objects.create(
            student=student,
            particular=particular,
            amount_paid=amount,
            receipt_no=receipt_no
        )
        messages.success(request, "Fee collected and receipt generated successfully!")
        return redirect('submit_fees')
    
    students = Student.objects.all()
    particulars = FeeParticular.objects.all()
    payments = FeePayment.objects.all().order_by('-payment_date')
    return render(request, 'core/submitfees.html', {'students': students, 'particulars': particulars, 'payments': payments})

def fee_particulars_view(request):
    if request.method == 'POST':
        title = request.POST.get('title')
        desc = request.POST.get('description', '')
        FeeParticular.objects.create(title=title, description=desc)
        messages.success(request, "Fee head added successfully!")
        return redirect('fee_particulars')
    particulars = FeeParticular.objects.all()
    return render(request, 'core/feeparticulars.html', {'particulars': particulars})

def mark_attendance_view(request):
    if request.method == 'POST':
        class_id = request.POST.get('class_section')
        date = request.POST.get('date', timezone.now().date())
        cls_obj = get_object_or_404(ClassSection, id=class_id)
        students = Student.objects.filter(class_section=cls_obj)
        
        for student in students:
            status = request.POST.get(f'status_{student.id}', 'Present')
            Attendance.objects.update_or_create(
                student=student,
                date=date,
                defaults={'status': status}
            )
        messages.success(request, "Attendance marked successfully!")
        return redirect('mark_attendance')
    
    classes = ClassSection.objects.all()
    selected_class_id = request.GET.get('class_section')
    students = []
    if selected_class_id:
        students = Student.objects.filter(class_section_id=selected_class_id)
    return render(request, 'core/markatt.html', {'classes': classes, 'students': students, 'selected_class_id': selected_class_id})

def employee_list_view(request):
    if request.method == 'POST':
        Employee.objects.create(
            employee_id=request.POST.get('employee_id'),
            name=request.POST.get('name'),
            designation=request.POST.get('designation'),
            phone=request.POST.get('phone'),
            salary=request.POST.get('salary', 0)
        )
        messages.success(request, "Employee registered successfully!")
        return redirect('employees')
    employees = Employee.objects.all()
    return render(request, 'core/empregfs.html', {'employees': employees})

def exam_marks_view(request):
    if request.method == 'POST':
        exam_id = request.POST.get('exam')
        subject_id = request.POST.get('subject')
        student_id = request.POST.get('student')
        marks = request.POST.get('marks_obtained')
        max_m = request.POST.get('max_marks', 100)
        
        exam = get_object_or_404(Exam, id=exam_id)
        subject = get_object_or_404(Subject, id=subject_id)
        student = get_object_or_404(Student, id=student_id)
        
        ExamMark.objects.update_or_create(
            exam=exam, student=student, subject=subject,
            defaults={'marks_obtained': marks, 'max_marks': max_m}
        )
        messages.success(request, "Exam marks saved successfully!")
        return redirect('exam_marks')
        
    exams = Exam.objects.all()
    subjects = Subject.objects.all()
    students = Student.objects.all()
    marks_list = ExamMark.objects.all()
    return render(request, 'core/examMarks.html', {'exams': exams, 'subjects': subjects, 'students': students, 'marks_list': marks_list})

def export_students_excel(request):
    response = HttpResponse(content_type='application/ms-excel')
    response['Content-Disposition'] = 'attachment; filename="students_list.xlsx"'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(['Registration No', 'Student Name', 'Class', 'Section', 'Father Name', 'Phone'])
    for s in Student.objects.all():
        ws.append([s.registration_no, s.student_name, s.class_section.class_name if s.class_section else '', s.section, s.father_name, s.phone_number])
    wb.save(response)
    return response

def import_students_excel(request):
    if request.method == 'POST' and request.FILES.get('excel_file'):
        file = request.FILES['excel_file']
        wb = openpyxl.load_workbook(file)
        sheet = wb.active
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0]:
                reg_no, name, c_name, sec, father, phone = row[0], row[1], row[2], row[3], row[4], row[5]
                cls_obj, _ = ClassSection.objects.get_or_create(class_name=str(c_name), section=str(sec) if sec else '')
                Student.objects.get_or_create(registration_no=str(reg_no), defaults={
                    'student_name': str(name),
                    'class_section': cls_obj,
                    'father_name': str(father) if father else '',
                    'phone_number': str(phone) if phone else ''
                })
        messages.success(request, "Students imported successfully via Excel!")
    return redirect('dashboard')